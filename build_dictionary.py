import os
import json
import re
import docx

def clean_text(text):
    if not text:
        return ""
    # Replace curly quotes and dashes correctly, and make sure we do NOT insert hyphens between chars!
    text = text.strip()
    text = text.replace("\u2013", "-").replace("\u2014", "-")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    text = text.replace("\u2018", "'").replace("\u2019", "'")
    return text

def parse_purpose_document(file_path):
    doc = docx.Document(file_path)
    filename = os.path.basename(file_path)
    
    sections = {}
    current_section = None
    
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        
        # Check for headings like A. PURPOSE DEFINITION
        match = re.match(r'^([A-F])\.\s+', text)
        if match:
            current_section = match.group(1)
            sections[current_section] = []
        elif current_section:
            sections[current_section].append(clean_text(text))
            
    # Parse the contents of sections
    purpose_def = ""
    if "A" in sections and sections["A"]:
        text_joined = " ".join(sections["A"])
        match = re.search(r'"Purpose"\s+means\s+(.+)', text_joined, re.IGNORECASE)
        if match:
            purpose_def = match.group(1).strip()
        else:
            purpose_def = text_joined
            
    recitals = []
    if "B" in sections:
        recitals = [r for r in sections["B"] if r.startswith("WHEREAS")]
        
    additional_defs = {}
    if "C" in sections:
        current_def_name = None
        for item in sections["C"]:
            if len(item) < 40 and not item.startswith('"') and not item.endswith('"'):
                current_def_name = item
                additional_defs[current_def_name] = ""
            elif current_def_name:
                additional_defs[current_def_name] = item
                
    ci_expansion = []
    if "D" in sections:
        for item in sections["D"]:
            if re.match(r'^\([a-z]\)\s+', item) or item.startswith("-") or item.strip().startswith("•"):
                ci_expansion.append(item)
            elif len(item) > 30 and "Confidential Information includes" not in item:
                ci_expansion.append(item)
                
    introduced_contacts = []
    if "E" in sections:
        for item in sections["E"]:
            if len(item) < 60 and "AI" not in item and "Trigger" not in item and "Replace" not in item and "Insert" not in item:
                introduced_contacts.append(item)
                
    return {
        "filename": filename,
        "purpose_definition": purpose_def,
        "recitals": recitals,
        "additional_definitions": additional_defs,
        "confidential_information_expansion": ci_expansion,
        "introduced_contacts": introduced_contacts
    }

def parse_master_clauses(file_path):
    doc = docx.Document(file_path)
    clauses = {}
    current_clause = None
    intro_paragraphs = []
    
    HEADERS_OF_INTEREST = [
        "DEFINITIONS",
        "EXCLUSIONS FROM CONFIDENTIAL INFORMATION",
        "CONFIDENTIALITY OBLIGATIONS",
        "PERMITTED DISCLOSURES",
        "NON-USE OBLIGATIONS",
        "COMPELLED DISCLOSURE",
        "RETURN OR DESTRUCTION OF CONFIDENTIAL INFORMATION",
        "INTELLECTUAL PROPERTY RIGHTS AND NO LICENSE",
        "DATA PROTECTION AND INFORMATION SECURITY",
        "NO WARRANTY",
        "TERM AND TERMINATION",
        "REMEDIES AND INJUNCTIVE RELIEF",
        "LIMITATION OF LIABILITY",
        "INDEMNITY",
        "NON-CIRCUMVENTION AND NON-SOLICITATION",
        "PUBLIC ANNOUNCEMENTS"
    ]
    
    headers_clean = [h.strip().upper() for h in HEADERS_OF_INTEREST]
    other_headings = [
        "NO OBLIGATION TO PROCEED",
        "GOVERNING LAW AND DISPUTE RESOLUTION",
        "MISCELLANEOUS",
        "ASSIGNMENT",
        "ENTIRE AGREEMENT",
        "AMENDMENTS",
        "WAIVER",
        "SEVERABILITY",
        "COUNTERPARTS AND ELECTRONIC SIGNATURES",
        "IN WITNESS WHEREOF"
    ]
    
    for p in doc.paragraphs:
        text = clean_text(p.text.strip())
        if not text:
            continue
            
        text_upper_clean = text.upper().strip()
        
        if text_upper_clean in headers_clean:
            idx = headers_clean.index(text_upper_clean)
            current_clause = f"{idx+1}_{text_upper_clean.replace(' ', '_')}"
            clauses[current_clause] = []
        elif text_upper_clean in other_headings:
            current_clause = None
        elif current_clause:
            clauses[current_clause].append(text)
        else:
            intro_paragraphs.append(text)
            
    return {
        "intro": intro_paragraphs,
        "clauses": clauses
    }

def restructure_master_nda(parsed_nda, doc_type):
    master_clauses = parsed_nda["clauses"]
    structured = {
        "document_type": doc_type,
        "1_DEFINITIONS": {},
        "2_EXCLUSIONS_FROM_CONFIDENTIAL_INFORMATION": {},
        "3_CONFIDENTIALITY_OBLIGATIONS": {},
        "4_PERMITTED_DISCLOSURES": {},
        "5_NON_USE_OBLIGATIONS": {},
        "6_COMPELLED_DISCLOSURE": {},
        "7_RETURN_OR_DESTRUCTION_OF_CONFIDENTIAL_INFORMATION": {},
        "8_INTELLECTUAL_PROPERTY_RIGHTS_AND_NO_LICENSE": {},
        "9_DATA_PROTECTION_AND_INFORMATION_SECURITY": {},
        "10_NO_WARRANTY": {},
        "11_TERM_AND_TERMINATION": {},
        "12_REMEDIES_AND_INJUNCTIVE_RELIEF": {},
        "13_LIMITATION_OF_LIABILITY": {},
        "14_INDEMNITY": {},
        "15_NON_CIRCUMVENTION_AND_NON_SOLICITATION": {},
        "16_PUBLIC_ANNOUNCEMENTS": {}
    }
    
    # 1. Parse Definitions
    defs_list = master_clauses.get("1_DEFINITIONS", [])
    current_term = None
    for idx, para in enumerate(defs_list):
        if para.startswith("For the purposes of this Agreement"):
            continue
        if (para.startswith('"') and para.endswith('"')) or len(para) < 30:
            current_term = para.replace('"', '').strip()
        elif current_term:
            key_num = len(structured["1_DEFINITIONS"]) + 1
            key_name = f"1.{key_num}_{current_term.replace(' ', '_')}"
            structured["1_DEFINITIONS"][key_name] = {
                "term": current_term,
                "definition": para
            }
            current_term = None
            
    # 2. Parse Exclusions
    exclusions_list = master_clauses.get("2_EXCLUSIONS_FROM_CONFIDENTIAL_INFORMATION", [])
    if exclusions_list:
        p0 = exclusions_list[0]
        parts = re.split(r'\(([a-h])\)\s+', p0)
        items = {}
        if len(parts) > 1:
            for i in range(1, len(parts), 2):
                letter = parts[i]
                content = parts[i+1].strip()
                if content.endswith(";"):
                    content = content[:-1]
                items[letter] = content
            structured["2_EXCLUSIONS_FROM_CONFIDENTIAL_INFORMATION"]["2.1"] = items
        else:
            structured["2_EXCLUSIONS_FROM_CONFIDENTIAL_INFORMATION"]["2.1"] = {"text": p0}
            
        if len(exclusions_list) > 1:
            structured["2_EXCLUSIONS_FROM_CONFIDENTIAL_INFORMATION"]["2.2"] = {"text": exclusions_list[1]}
            
    # Numbered clauses helper
    def parse_numbered(clause_name, structured_key):
        clause_items = master_clauses.get(clause_name, [])
        for idx, para in enumerate(clause_items):
            key = f"{structured_key.split('_')[0]}.{idx+1}"
            structured[structured_key][key] = para

    parse_numbered("3_CONFIDENTIALITY_OBLIGATIONS", "3_CONFIDENTIALITY_OBLIGATIONS")
    parse_numbered("4_PERMITTED_DISCLOSURES", "4_PERMITTED_DISCLOSURES")
    parse_numbered("5_NON_USE_OBLIGATIONS", "5_NON-USE_OBLIGATIONS")
    parse_numbered("6_COMPELLED_DISCLOSURE", "6_COMPELLED_DISCLOSURE")
    parse_numbered("7_RETURN_OR_DESTRUCTION_OF_CONFIDENTIAL_INFORMATION", "7_RETURN_OR_DESTRUCTION_OF_CONFIDENTIAL_INFORMATION")
    parse_numbered("8_INTELLECTUAL_PROPERTY_RIGHTS_AND_NO_LICENSE", "8_INTELLECTUAL_PROPERTY_RIGHTS_AND_NO_LICENSE")
    parse_numbered("9_DATA_PROTECTION_AND_INFORMATION_SECURITY", "9_DATA_PROTECTION_AND_INFORMATION_SECURITY")
    parse_numbered("10_NO_WARRANTY", "10_NO_WARRANTY")
    parse_numbered("11_TERM_AND_TERMINATION", "11_TERM_AND_TERMINATION")
    parse_numbered("12_REMEDIES_AND_INJUNCTIVE_RELIEF", "12_REMEDIES_AND_INJUNCTIVE_RELIEF")
    parse_numbered("13_LIMITATION_OF_LIABILITY", "13_LIMITATION_OF_LIABILITY")
    parse_numbered("14_INDEMNITY", "14_INDEMNITY")
    parse_numbered("15_NON_CIRCUMVENTION_AND_NON_SOLICITATION", "15_NON_CIRCUMVENTION_AND_NON_SOLICITATION")
    parse_numbered("16_PUBLIC_ANNOUNCEMENTS", "16_PUBLIC_ANNOUNCEMENTS")
    
    return structured

def merge_purpose_into_structured(master_struct, purpose_data, document_name):
    # Create a deep copy of the structured clauses
    res = json.loads(json.dumps(master_struct))
    res["document_type"] = f"{master_struct['document_type']} - {document_name}"
    
    # 1. Replace Purpose Definition
    if purpose_data["purpose_definition"]:
        purpose_val = purpose_data["purpose_definition"]
        # Find Purpose key in Definitions
        purpose_key = None
        for k, v in res["1_DEFINITIONS"].items():
            if v["term"].lower() == "purpose":
                purpose_key = k
                break
        if purpose_key:
            res["1_DEFINITIONS"][purpose_key]["definition"] = f'"Purpose" means {purpose_val}'
            
    # 2. Expand Confidential Information
    if purpose_data["confidential_information_expansion"]:
        ci_exp_str = " " + " ".join(purpose_data["confidential_information_expansion"])
        # Find Confidential Information key
        ci_key = None
        for k, v in res["1_DEFINITIONS"].items():
            if "confidential" in v["term"].lower() and "information" in v["term"].lower():
                ci_key = k
                break
        if ci_key:
            res["1_DEFINITIONS"][ci_key]["definition"] += ci_exp_str
            
    # 3. Add Additional Definitions
    for term, def_val in purpose_data["additional_definitions"].items():
        key_num = len(res["1_DEFINITIONS"]) + 1
        key_name = f"1.{key_num}_{term.replace(' ', '_')}"
        res["1_DEFINITIONS"][key_name] = {
            "term": term,
            "definition": def_val
        }
        
    return res

def group_by_clause_type(doc_list):
    grouped = {
        "1_DEFINITIONS": [],
        "2_EXCLUSIONS_FROM_CONFIDENTIAL_INFORMATION": [],
        "3_CONFIDENTIALITY_OBLIGATIONS": [],
        "4_PERMITTED_DISCLOSURES": [],
        "5_NON_USE_OBLIGATIONS": [],
        "6_COMPELLED_DISCLOSURE": [],
        "7_RETURN_OR_DESTRUCTION_OF_CONFIDENTIAL_INFORMATION": [],
        "8_INTELLECTUAL_PROPERTY_RIGHTS_AND_NO_LICENSE": [],
        "9_DATA_PROTECTION_AND_INFORMATION_SECURITY": [],
        "10_NO_WARRANTY": [],
        "11_TERM_AND_TERMINATION": [],
        "12_REMEDIES_AND_INJUNCTIVE_RELIEF": [],
        "13_LIMITATION_OF_LIABILITY": [],
        "14_INDEMNITY": [],
        "15_NON_CIRCUMVENTION_AND_NON_SOLICITATION": [],
        "16_PUBLIC_ANNOUNCEMENTS": []
    }
    
    for doc in doc_list:
        doc_type = doc["document_type"]
        for key in grouped.keys():
            if key in doc:
                grouped[key].append({
                    "document_type": doc_type,
                    "clauses": doc[key]
                })
    return grouped

def main():
    docs_dir = r"d:\ai agent\new_legal_ai\docs"
    master_mutual_file = os.path.join(docs_dir, "Master NDA Zip File", "Master NDA", "NDA_01_Master_Mutual_NDA_Benchmark_v1.0.docx")
    master_unilateral_file = os.path.join(docs_dir, "Master NDA Zip File", "Master NDA", "UNILATERAL NDA v1.0.docx")
    
    # 1. Parse Masters
    parsed_mutual = parse_master_clauses(master_mutual_file)
    parsed_unilateral = parse_master_clauses(master_unilateral_file)
    
    # 2. Restructure Masters
    structured_mutual = restructure_master_nda(parsed_mutual, "Mutual NDA")
    structured_unilateral = restructure_master_nda(parsed_unilateral, "Unilateral NDA")
    
    # Save structured master files to backend/data
    out_data_dir = r"d:\ai agent\new_legal_ai\backend\data"
    os.makedirs(out_data_dir, exist_ok=True)
    
    with open(os.path.join(out_data_dir, "mutual_nda_structured.json"), "w", encoding="utf-8") as f:
        json.dump(structured_mutual, f, indent=2, ensure_ascii=False)
    with open(os.path.join(out_data_dir, "unilateral_nda_structured.json"), "w", encoding="utf-8") as f:
        json.dump(structured_unilateral, f, indent=2, ensure_ascii=False)
        
    print("Structured Master NDAs saved to backend/data")
    
    # Save clean raw master json as well
    with open(os.path.join(out_data_dir, "mutual_nda_master.json"), "w", encoding="utf-8") as f:
        json.dump(parsed_mutual, f, indent=2, ensure_ascii=False)
    with open(os.path.join(out_data_dir, "unilateral_nda_master.json"), "w", encoding="utf-8") as f:
        json.dump(parsed_unilateral, f, indent=2, ensure_ascii=False)
        
    # 3. Parse Purpose files (GC and RE)
    gc_dir = os.path.join(docs_dir, "General Commercial NDA Zip File", "General Commercial NDA")
    re_dir = os.path.join(docs_dir, "Renewable Energy NDA database Zip File", "Renewable Energy NDA database")
    
    compiled_database = {
        "General Commercial": {},
        "Renewable Energy": {}
    }
    
    all_docs_structured = []
    gc_docs_structured = []
    re_docs_structured = []
    
    # GC
    if os.path.exists(gc_dir):
        for f in sorted(os.listdir(gc_dir)):
            if f.endswith(".docx") and not f.startswith("~$"):
                try:
                    res = parse_purpose_document(os.path.join(gc_dir, f))
                    # Clean title
                    title = f.replace("GC-", "").split(".docx")[0]
                    title = re.sub(r'^\d+[\s\-_]*', '', title)
                    title = title.replace("Discussions", "").replace("Discussion", "").replace("DICUSSION", "").strip()
                    title = " ".join([w.capitalize() for w in title.split()])
                    if not title:
                        title = f
                    
                    compiled_database["General Commercial"][title] = res
                    
                    # Merge to generate fully structured documents
                    gc_mutual = merge_purpose_into_structured(structured_mutual, res, f"General Commercial {title} (Mutual)")
                    gc_unilateral = merge_purpose_into_structured(structured_unilateral, res, f"General Commercial {title} (Unilateral)")
                    
                    gc_docs_structured.extend([gc_mutual, gc_unilateral])
                    all_docs_structured.extend([gc_mutual, gc_unilateral])
                except Exception as e:
                    print(f"Error parsing GC file {f}: {e}")
                    
    # RE
    if os.path.exists(re_dir):
        for f in sorted(os.listdir(re_dir)):
            if (f.endswith(".docx") or f.endswith(".doc.docx")) and not f.startswith("~$"):
                try:
                    res = parse_purpose_document(os.path.join(re_dir, f))
                    title = f.replace("RE-", "").split(".docx")[0].split(".doc")[0]
                    title = re.sub(r'^\d+[\s\-_]*', '', title)
                    title = title.replace("Discussions", "").replace("Discussion", "").strip()
                    title = " ".join([w.capitalize() for w in title.split()])
                    if not title:
                        title = f
                        
                    compiled_database["Renewable Energy"][title] = res
                    
                    # Merge
                    re_mutual = merge_purpose_into_structured(structured_mutual, res, f"Renewable Energy {title} (Mutual)")
                    re_unilateral = merge_purpose_into_structured(structured_unilateral, res, f"Renewable Energy {title} (Unilateral)")
                    
                    re_docs_structured.extend([re_mutual, re_unilateral])
                    all_docs_structured.extend([re_mutual, re_unilateral])
                except Exception as e:
                    print(f"Error parsing RE file {f}: {e}")
                    
    # Save clean nda_purposes.json in backend/data
    with open(os.path.join(out_data_dir, "nda_purposes.json"), "w", encoding="utf-8") as f:
        json.dump(compiled_database, f, indent=2, ensure_ascii=False)
        
    print("Compiled purposes (nda_purposes.json) saved successfully with clean text (no hyphenation bug).")
    
    # 4. Create Dictionary directory and write grouped clauses
    dict_dir = r"d:\ai agent\new_legal_ai\dictionary"
    os.makedirs(dict_dir, exist_ok=True)
    
    # Master Grouped
    master_grouped = group_by_clause_type([structured_mutual, structured_unilateral])
    with open(os.path.join(dict_dir, "master_ndas.json"), "w", encoding="utf-8") as f:
        json.dump(master_grouped, f, indent=2, ensure_ascii=False)
        
    # General Commercial Grouped
    gc_grouped = group_by_clause_type(gc_docs_structured)
    with open(os.path.join(dict_dir, "general_commercial.json"), "w", encoding="utf-8") as f:
        json.dump(gc_grouped, f, indent=2, ensure_ascii=False)
        
    # Renewable Energy Grouped
    re_grouped = group_by_clause_type(re_docs_structured)
    with open(os.path.join(dict_dir, "renewable_energy.json"), "w", encoding="utf-8") as f:
        json.dump(re_grouped, f, indent=2, ensure_ascii=False)
        
    # All Grouped
    all_grouped = group_by_clause_type([structured_mutual, structured_unilateral] + all_docs_structured)
    with open(os.path.join(dict_dir, "all_grouped.json"), "w", encoding="utf-8") as f:
        json.dump(all_grouped, f, indent=2, ensure_ascii=False)
        
    print(f"Successfully compiled all clauses into dictionary folder at {dict_dir}")
    print("All tasks completed.")

if __name__ == "__main__":
    main()
