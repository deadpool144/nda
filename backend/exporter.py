import os
import io
import re
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def export_to_docx(title, text):
    doc = docx.Document()
    
    # Configure page setup
    section = doc.sections[0]
    section.page_width = Inches(8.27)  # A4 width
    section.page_height = Inches(11.69) # A4 height
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Set default style to Times New Roman, 12pt, Justified, Single spacing
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(title.upper())
    run_title.bold = True
    run_title.size = Pt(14)
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)
    
    # Split the contract text by paragraph and write
    lines = text.split("\n")
    skip_title = True
    
    for line in lines:
        line_strip = line.strip()
        if not line_strip:
            continue
            
        # Skip the title in the body if it matches the main title
        if skip_title and line_strip.upper() == title.upper():
            skip_title = False
            continue
            
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        # Heading Check
        is_heading, heading_level = get_heading_info(line_strip)
        
        if is_heading:
            run = p.add_run(line_strip)
            run.bold = True
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if heading_level == 1:
                run.size = Pt(12)
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
            elif heading_level == 2:
                run.size = Pt(12)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(3)
            elif heading_level == 3: # Special centered headers like RECITALS
                run.size = Pt(12)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
        else:
            # Handle list indentation
            if re.match(r'^\([a-zivx]+\)', line_strip):
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.25)
            elif re.match(r'^[\d\.]+\s', line_strip) and not is_heading:
                # Sub-clause text or definitions body indenting can be handled here if needed
                p.paragraph_format.left_indent = Inches(0)
                
            # Check for defined terms bolding
            definition_match = re.match(r'^(\d+\.\d+(?:\.\d+)?)\s+("[^"]+")', line_strip)
            if definition_match:
                num = definition_match.group(1)
                term = definition_match.group(2)
                remainder = line_strip[len(num) + len(term) + 1:].strip()
                
                run_num = p.add_run(num + " ")
                run_term = p.add_run(term)
                run_term.bold = True
                
                if remainder:
                    p.add_run(" " + remainder)
            else:
                p.add_run(line_strip)
                
    # Save to a memory stream
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

def get_heading_info(text):
    # Returns (is_heading, level)
    if text in ["RECITALS", "BY AND BETWEEN", "AND", "IN WITNESS WHEREOF"]:
        return True, 3 # Center aligned headers
    
    # Level 1: "1. DEFINITIONS", "10. NOTICE DETAILS"
    if re.match(r'^(\d+)\.\s+[A-Z\s\-]+$', text):
        return True, 1
        
    # Level 2: "1.1 Governing Law", "1.2 Arbitration"
    if re.match(r'^(\d+\.\d+)\s+[A-Za-z\s\-]+$', text):
        # We only bold subheadings if they are relatively short (titles)
        if len(text.split()) < 8 and not text.endswith('.'):
            return True, 2
            
    # All caps generic section headers
    if re.match(r'^[A-Z][A-Z\s]{4,30}$', text) and len(text) < 40:
        return True, 1
        
    return False, 0
